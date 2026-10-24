"""Build the hackathon submission (Word) from SUBMISSION.md, inside the Apart Research template.

The template's title page is kept: the title, a single author line, "With Apart Research", the sprint footnote and
the abstract box are filled in; its guidance text is replaced by the report body. pandoc converts the body (with the
template as reference document, so the template's heading styles are used; equations become Word equations), and
docxcompose appends it to the filled template. Then fonts, tables, captions and page numbers get the template's look.

Usage: python report/submission/make_submission.py  -> report/submission/Apart_submission.docx
"""
import copy
import io
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from docxcompose.composer import Composer

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import make_docx as M  # noqa: E402  (shared helpers: element builder, fonts, table widths, schema order)

SRC, TEMPLATE, OUT = HERE / "SUBMISSION.md", HERE / "template.docx", HERE / "Apart_submission.docx"
FONT, MONO = "Old Standard TT", "Courier New"
AUTHOR = "Shantanu P. Sontakke"
TEXT_W_DXA = 9360  # US Letter, 1-inch margins (as in the template)


def reference_without_fonts(dest: Path) -> Path:
    """pandoc copies a reference document's font list but not its embedded fonts, which leaves a broken package;
    give pandoc a copy of the template without them (the final document is built on the original template)."""
    with zipfile.ZipFile(TEMPLATE) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/fonts/"):
                continue
            if item.filename == "word/_rels/fontTable.xml.rels":
                data = re.sub(rb"<Relationship [^>]*/>", b"", data)
            if item.filename == "word/fontTable.xml":
                data = re.sub(rb"<w:embed(Regular|Bold|Italic|BoldItalic)[^>]*/>", b"", data)
            zout.writestr(item, data)
    return dest


def parts() -> tuple[str, str, str]:
    text = SRC.read_text()
    title = re.search(r'^title:\s*"(.*)"\s*$', text, re.M).group(1)
    abstract = re.search(r"<!-- ABSTRACT -->\n(.*?)\n<!-- /ABSTRACT -->", text, re.S).group(1).strip()
    body = text.split("<!-- /ABSTRACT -->", 1)[1]
    return title, abstract, body


def run_xml(text: str, size_half_pts: int | None = None, bold: bool = False) -> "M.OxmlElement":
    r = M.el("w:r")
    rpr = M.el("w:rPr")
    M.set_fonts(rpr, FONT)
    if bold:
        rpr.append(M.el("w:b"))
    if size_half_pts:
        rpr.append(M.el("w:sz", val=size_half_pts))
        rpr.append(M.el("w:szCs", val=size_half_pts))
    r.append(rpr)
    t = M.el("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def fill_title_page(doc: Document, title: str, abstract: str) -> None:
    front = doc.element.body[1]
    for t in front.iter(qn("w:t")):
        if t.text == "PROJECT TITLE":
            t.text = title
    # the six-slot author grid becomes one centred author line
    grids = [tb for tb in front.iter(qn("w:tbl")) if "Author name 1" in "".join(x.text or "" for x in tb.iter(qn("w:t")))
             and tb.find(".//" + qn("w:tbl")) is None]
    grid = grids[0]
    p = M.el("w:p")
    ppr = M.el("w:pPr")
    ppr.append(M.el("w:spacing", after=120))
    ppr.append(M.el("w:jc", val="center"))
    p.append(ppr)
    p.append(run_xml(AUTHOR, 24))
    grid.addprevious(p)
    grid.getparent().remove(grid)
    # the abstract box: guidance text replaced by the abstract, plain (not italic or underlined)
    for para in front.iter(qn("w:p")):
        txt = "".join(x.text or "" for x in para.iter(qn("w:t")))
        if txt.startswith("Summarize your project"):
            for r in list(para.findall(qn("w:r"))):
                para.remove(r)
            ppr = para.find(qn("w:pPr"))
            rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
            if rpr is not None:
                ppr.remove(rpr)
            para.append(run_xml(abstract))
    # everything after the title page (info box and guidance) goes; the section settings stay
    body = doc.element.body
    for child in list(body)[2:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def copy_missing_styles(dst: Document) -> None:
    """With this template as reference document pandoc keeps only the template's styles, so the body's paragraph and
    table styles (Body Text, Compact, Table, captions, ...) are undefined; add pandoc's defaults for every style the
    merged document lacks (style_body then gives them the template's look)."""
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True,
                          check=True).stdout
    src = Document(io.BytesIO(data))
    have = {s.get(qn("w:styleId")) for s in dst.styles.element.findall(qn("w:style"))}
    for st in src.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) not in have:
            dst.styles.element.append(copy.deepcopy(st))


def style_body(doc: Document) -> None:
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    M.set_fonts(defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}"), FONT)
    black = M.RGBColor(0, 0, 0)
    for name, size, bold, before, after, line in (("Normal", 11, None, None, None, None),
                                                  ("Body Text", 11, None, 0, 6, 1.15),
                                                  ("First Paragraph", 11, None, 0, 6, 1.15),
                                                  ("Compact", 11, None, 0, 3, 1.15),
                                                  ("Heading 2", 14, True, 12, 3, 1.15),
                                                  ("Heading 3", 12, True, 8, 2, 1.15)):
        try:
            M.style(doc, name, FONT, size, bold=bold, italic=False, color=black, before=before, after=after,
                    line=line, keep_next=True if name.startswith("Heading") else None)
        except KeyError:
            pass
    for name in ("Image Caption", "Table Caption"):
        try:
            M.style(doc, name, FONT, 9.5, italic=False, color=black, before=2, after=8, line=1.0)
        except KeyError:
            pass
    try:
        M.style(doc, "Verbatim Char", MONO, 9)
    except KeyError:
        pass


def format_tables(doc: Document) -> None:
    M.TEXT_W_DXA = TEXT_W_DXA
    for table in doc.tables[1:]:  # the first table is the template's title page
        M.format_table(table)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        M.set_fonts(r._r.get_or_add_rPr(), FONT)
                        r.font.size = Pt(9)


def tables_on_one_page(doc: Document) -> None:
    """Keep each table on one page: every row but the last stays with the next one."""
    for tbl in doc.tables[1:]:
        for row in tbl.rows[:-1]:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True


def space_after_tables(doc: Document) -> None:
    """Leave a gap between a table and whatever follows it (the body styles have no space before)."""
    for tbl in doc.tables[1:]:
        nxt = tbl._tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            Paragraph(nxt, doc._body).paragraph_format.space_before = Pt(8)


def keep_together(doc: Document) -> None:
    for p in doc.paragraphs:
        if re.match(r"\*?\*?Table [A-Z]?\d+\.", p.text) or p.text.startswith("Table "):
            if p.runs and p.runs[0].bold:
                p.paragraph_format.keep_with_next = True
        if p._p.findall(".//" + qn("w:drawing")):
            p.paragraph_format.keep_with_next = True
            for d in p._p.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"):
                cx, cy = int(d.get("cx")), int(d.get("cy"))
                limit = int(6.4 * 914400)
                if cx > limit:
                    d.set("cx", str(limit))
                    d.set("cy", str(int(cy * limit / cx)))


def page_numbers(doc: Document) -> None:
    for sec in doc.sections:
        sec.footer.is_linked_to_previous = False
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = M.el("w:fldSimple", instr="PAGE")
        fld.append(run_xml("1", 18))
        p._p.append(fld)


def main() -> int:
    title, abstract, body_md = parts()
    words = len(abstract.split())
    with tempfile.TemporaryDirectory() as tmp:
        body_docx = Path(tmp) / "body.docx"
        ref = reference_without_fonts(Path(tmp) / "reference.docx")
        subprocess.run(["pandoc", "-f", "markdown", "-o", str(body_docx), "--reference-doc", str(ref),
                        "--resource-path", str(HERE.parent)], input=body_md, text=True, check=True, cwd=HERE.parent)
        master = Document(str(TEMPLATE))
        fill_title_page(master, title, abstract)
        body_doc = Document(str(body_docx))
        composer = Composer(master)
        composer.append(body_doc)
        composer.save(str(OUT))
    doc = Document(str(OUT))
    copy_missing_styles(doc)
    style_body(doc)
    format_tables(doc)
    space_after_tables(doc)
    tables_on_one_page(doc)
    keep_together(doc)
    page_numbers(doc)
    M.fix_schema(doc)
    doc.core_properties.title = title
    doc.core_properties.author = AUTHOR
    doc.save(str(OUT))
    print(f"wrote {OUT} (abstract: {words} words)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
