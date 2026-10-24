"""Word version of the report: pandoc converts report/REPORT.md (equations become Word equations), then
python-docx gives it the same look as REPORT.pdf (A4, serif body, sans headings, ruled tables, page numbers).
Usage: python report/make_docx.py  -> report/REPORT.docx"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "REPORT.docx"
SERIF, SANS, MONO = "Georgia", "Arial", "Consolas"
INK, INK2, MUTED = RGBColor(0x11, 0x11, 0x11), RGBColor(0x33, 0x33, 0x33), RGBColor(0x66, 0x66, 0x66)
RULE, CELL_RULE, HEAD_FILL = "BBBBBB", "CCCCCC", "F1F1EE"
MARGIN_TB, MARGIN_LR = Mm(17), Mm(16)
TEXT_W_DXA = round((210 - 32) / 25.4 * 1440)  # A4 width minus the side margins, in twentieths of a point

TBL_PR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
                "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
                "tblCaption", "tblDescription"]
P_PR_AFTER_PBDR = ["w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                   "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                   "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                   "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                   "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]


def el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def set_fonts(rpr, name: str) -> None:
    """Explicit fonts for every script, dropping theme fonts (which would otherwise win)."""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = el("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{a}"), None)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{a}"), name)


def style(doc, name, font=None, size=None, bold=None, italic=None, color=None, before=None, after=None,
          line=None, align=None, keep_next=None):
    # match the style id: pandoc writes "Heading 1" where python-docx's name lookup expects "heading 1"
    s = next((x for x in doc.styles if x.style_id == name.replace(" ", "")), None)
    if s is None:
        raise KeyError(f"style not found: {name}")
    if font:
        set_fonts(s.element.get_or_add_rPr(), font)
    f = s.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color
        c = s.element.get_or_add_rPr().find(qn("w:color"))
        for a in ("themeColor", "themeTint", "themeShade"):  # theme colours would override the explicit one
            c.attrib.pop(qn(f"w:{a}"), None)
    pf = getattr(s, "paragraph_format", None)
    if pf is not None:
        if before is not None:
            pf.space_before = Pt(before)
        if after is not None:
            pf.space_after = Pt(after)
        if line is not None:
            pf.line_spacing = line
        if align is not None:
            pf.alignment = align
        if keep_next is not None:
            pf.keep_with_next = keep_next
    return s


def restyle(doc) -> None:
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr = defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
        if rpr is not None:
            set_fonts(rpr, SERIF)
    style(doc, "Normal", SERIF, 10.5, color=INK)
    for name in ("Body Text", "First Paragraph"):
        style(doc, name, SERIF, 10.5, color=INK, before=0, after=6, line=1.15)
    style(doc, "Compact", SERIF, 10.5, color=INK, before=0, after=3, line=1.15)
    style(doc, "Title", SANS, 20, bold=True, color=INK, before=0, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    style(doc, "Subtitle", SANS, 12.5, bold=False, italic=False, color=INK2, before=0, after=10,
          align=WD_ALIGN_PARAGRAPH.LEFT)
    h1 = style(doc, "Heading 1", SANS, 13, bold=True, italic=False, color=INK, before=16, after=5, keep_next=True)
    style(doc, "Heading 2", SANS, 11, bold=True, italic=False, color=INK, before=11, after=3, keep_next=True)
    style(doc, "Image Caption", SANS, 8.5, italic=False, color=INK2, before=2, after=10)
    style(doc, "Verbatim Char", MONO, 8.5)
    if h1 is not None:  # thin rule under section headings, as in the PDF
        ppr = h1.element.get_or_add_pPr()
        old = ppr.find(qn("w:pBdr"))
        if old is not None:
            ppr.remove(old)
        bdr = el("w:pBdr")
        bdr.append(el("w:bottom", val="single", sz=6, space=2, color=RULE))
        ppr.insert_element_before(bdr, *P_PR_AFTER_PBDR)


def column_widths(table, total: int) -> list[int]:
    """Every column fits its longest word (no word is split); the rest of the width follows the text length."""
    char, pad = 100, 230  # about one 8.5 pt character (bold-safe) and the cell's side margins, in DXA
    texts = [[c.text.strip() for c in row.cells] for row in table.rows]
    least, want = [], []
    for j in range(len(table.columns)):
        col = [r[j] for r in texts]
        word = max((len(w) for t in col for w in t.split()), default=1)
        body = max((len(t) for t in col[1:]), default=0)
        least.append(word * char + pad)
        want.append(max(word, min(body, 30), min(len(col[0]), 12)) * char + pad)
    if sum(want) <= total:
        widths = [int(total * w / sum(want)) for w in want]
    else:
        extra = [w - m for w, m in zip(want, least)]
        room = max(total - sum(least), 0)
        widths = [m + int(room * e / sum(extra)) if sum(extra) else m for m, e in zip(least, extra)]
    widths[-1] += total - sum(widths)
    return widths


def format_table(table) -> None:
    widths = column_widths(table, TEXT_W_DXA)
    tbl = table._tbl
    tblPr = tbl.tblPr
    parts = {child.tag.split("}")[1]: child for child in list(tblPr)}
    for child in list(tblPr):
        tblPr.remove(child)
    parts["tblW"] = el("w:tblW", w=sum(widths), type="dxa")
    borders = el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in ("left", "right", "insideV"):
            borders.append(el(f"w:{edge}", val="nil"))
        else:
            borders.append(el(f"w:{edge}", val="single", sz=4, space=0, color=CELL_RULE))
    parts["tblBorders"] = borders
    parts["tblLayout"] = el("w:tblLayout", type="fixed")
    mar = el("w:tblCellMar")
    for side, v in (("top", 45), ("left", 85), ("bottom", 45), ("right", 85)):
        mar.append(el(f"w:{side}", w=v, type="dxa"))
    parts["tblCellMar"] = mar
    for tag in TBL_PR_ORDER:
        if tag in parts:
            tblPr.append(parts[tag])
    grid = tbl.tblGrid
    for g in list(grid):
        grid.remove(g)
    for w in widths:
        grid.append(el("w:gridCol", w=w))
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        for child in list(trPr):
            if child.tag in (qn("w:cantSplit"), qn("w:tblHeader")):
                trPr.remove(child)
        trPr.append(el("w:cantSplit"))
        if i == 0:
            trPr.append(el("w:tblHeader"))
        for j, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            for child in list(tcPr):
                if child.tag in (qn("w:tcW"), qn("w:shd")):
                    tcPr.remove(child)
            tcPr.insert(0, el("w:tcW", w=widths[j], type="dxa"))
            if i == 0:
                tcPr.insert(1, el("w:shd", val="clear", color="auto", fill=HEAD_FILL))
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0
                for r in p.runs:
                    set_fonts(r._r.get_or_add_rPr(), SANS)
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = INK
                    if i == 0:
                        r.font.bold = True


def keep_captions_with_content(doc) -> None:
    for p in doc.paragraphs:
        if p.text.startswith("Table ") and p.runs and p.runs[0].bold:
            p.paragraph_format.keep_with_next = True  # table titles stay with their table
        if p._p.findall(f".//{qn('w:drawing')}"):
            p.paragraph_format.keep_with_next = True  # figures stay with their caption


def page_setup(doc) -> None:
    for sec in doc.sections:
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.top_margin = sec.bottom_margin = MARGIN_TB
        sec.left_margin = sec.right_margin = MARGIN_LR
        sec.footer.is_linked_to_previous = False
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = el("w:fldSimple", instr="PAGE")
        run = el("w:r")
        rpr = el("w:rPr")
        set_fonts(rpr, SANS)
        rpr.append(el("w:color", val=str(MUTED)))
        rpr.append(el("w:sz", val=17))
        run.append(rpr)
        t = el("w:t")
        t.text = "1"
        run.append(t)
        fld.append(run)
        p._p.append(fld)


# Child order required by the Office Open XML schema; pandoc's output and python-docx edits do not always keep it.
SEQ = {
    "w:style": "name aliases basedOn next link autoRedefine hidden uiPriority semiHidden unhideWhenUsed qFormat locked "
               "personal personalCompose personalReply rsid pPr rPr tblPr trPr tcPr tblStylePr",
    "w:pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
             "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
             "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
             "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "w:rPr": "ins del moveFrom moveTo rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss "
             "imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr "
             "shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath rPrChange",
    "w:tblPr": " ".join(TBL_PR_ORDER) + " tblPrChange",
    "w:trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden "
              "ins del trPrChange",
    "w:tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark "
              "headers cellIns cellDel cellMerge tcPrChange",
    "m:dPr": "begChr sepChr endChr grow shp ctrlPr",
}


SETTINGS_ORDER = " ".join(f"w:{t}" if ":" not in t else t for t in (
    "writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries "
    "displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData embedTrueTypeFonts "
    "embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins alignBordersAndEdges bordersDoNotSurroundHeader "
    "bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors hideGrammaticalErrors activeWritingStyle proofState "
    "formsDesign attachedTemplate linkStyles stylePaneFormatFilter stylePaneSortMethod documentType mailMerge "
    "revisionView trackRevisions doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride "
    "styleLockTheme styleLockQFSet defaultTabStop autoHyphenation consecutiveHyphenLimit hyphenationZone "
    "doNotHyphenateCaps showEnvelope summaryLength clickAndTypeStyle defaultTableStyle evenAndOddHeaders "
    "bookFoldRevPrinting bookFoldPrinting bookFoldPrintingSheets drawingGridHorizontalSpacing "
    "drawingGridVerticalSpacing displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery "
    "doNotUseMarginsForDrawingGridOrigin drawingGridHorizontalOrigin drawingGridVerticalOrigin doNotShadeFormData "
    "noPunctuationKerning characterSpacingControl printTwoOnOne strictFirstAndLastChars noLineBreaksAfter "
    "noLineBreaksBefore savePreviewPicture doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent "
    "alwaysShowPlaceholderText doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt "
    "showXMLTags alwaysMergeEmptyNamespace updateFields hdrShapeDefaults footnotePr endnotePr compat docVars rsids "
    "m:mathPr attachedSchema themeFontLang clrSchemeMapping doNotIncludeSubdocsInStats doNotAutoCompressPictures "
    "forceUpgrade captions readModeInkLockDown smartTagType schemaLibrary shapeDefaults doNotEmbedSmartTags "
    "decimalSymbol listSeparator").split())


def reorder(parent, seq: str, prefix: str) -> None:
    rank = {qn(f"{prefix}:{t}"): i for i, t in enumerate(seq.split())}
    kids = [k for k in parent if isinstance(k.tag, str)]
    ordered = sorted(kids, key=lambda k: rank.get(k.tag, len(rank)))  # stable: unknown tags keep their order, last
    if ordered != kids:
        for k in kids:
            parent.remove(k)
        for k in ordered:
            parent.append(k)


def fix_schema(doc) -> None:
    roots = [doc.element, doc.styles.element, doc.part.numbering_part.element]
    roots += [sec.footer._element for sec in doc.sections]
    for root in roots:
        for tag, seq in SEQ.items():
            prefix, name = tag.split(":")
            for node in root.iter(qn(tag)):
                reorder(node, seq, prefix)
        for mrpr in root.iter(qn("m:rPr")):  # 'normal text' and a math style are alternatives, not both
            if mrpr.find(qn("m:nor")) is not None:
                for extra in mrpr.findall(qn("m:sty")) + mrpr.findall(qn("m:scr")):
                    mrpr.remove(extra)
    for node in doc.part.numbering_part.element.iter(qn("w:nsid"), qn("w:tmpl")):
        node.set(qn("w:val"), node.get(qn("w:val")).zfill(8).upper())  # 4-byte hex numbers
    for pgmar in doc.element.iter(qn("w:pgMar")):
        for attr, default in (("header", 708), ("footer", 708), ("gutter", 0)):
            if pgmar.get(qn(f"w:{attr}")) is None:
                pgmar.set(qn(f"w:{attr}"), str(default))
    for rpr in doc.styles.element.iter(qn("w:rPr")):  # pandoc's AbstractTitle style carries a stray '>' character
        rpr.text = None
        for child in rpr:
            child.tail = None
    settings = doc.settings.element
    rank = {qn(t): i for i, t in enumerate(SETTINGS_ORDER.split())}
    kids = [k for k in settings if isinstance(k.tag, str)]
    for k in kids:
        settings.remove(k)
    for k in sorted(kids, key=lambda k: rank.get(k.tag, len(rank))):
        settings.append(k)


def main() -> int:
    subprocess.run(["pandoc", "REPORT.md", "--shift-heading-level-by=-1", "-o", OUT.name], cwd=HERE, check=True)
    doc = Document(str(OUT))
    restyle(doc)
    for table in doc.tables:
        format_table(table)
    keep_captions_with_content(doc)
    page_setup(doc)
    fix_schema(doc)
    doc.core_properties.title = "Hints Win Far More Agreement Than Explicit Collusion Deals"
    doc.core_properties.subject = "Apart Research AI Collusion Sprint, research question 3.2"
    doc.save(str(OUT))
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
